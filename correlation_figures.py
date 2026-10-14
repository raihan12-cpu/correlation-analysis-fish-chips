# ============================================================
# USER SETTINGS / EDIT HERE
# ============================================================
from pathlib import Path
ROOT = Path(__file__).resolve().parent
DATA_FILE = ROOT / 'source' / 'Fish_Crackers_Correlation_PCA_HCA_Preprocessed.xlsx'
OUTPUT_DIR = ROOT / 'results'
FIGURE_DPI = 350
VARIABLES = ['Protein_pct', 'Ash_pct', 'Fat_pct', 'Moisture_pct', 'Linear_expansion_pct', 'Reported_oil_absorption_pct', 'Specific_volume_cm3_g', 'Color_L', 'Color_a', 'Color_b']
import sys, json, math, itertools
import numpy as np
import openpyxl
from correlation_core import require, worksheet_records, keyed, correlations as calculate_correlations
wb = openpyxl.load_workbook(DATA_FILE, data_only=True)
required = ['Sample_ID', 'Batch_ID', 'Batch_number', 'Fish_species', 'Fish_inclusion_pct', 'Is_control', *VARIABLES, 'Carbohydrate_pct']
data = worksheet_records(wb['Batch_data'], required)
keyed(data, ['Sample_ID', 'Batch_number'])
keys = [(r['Sample_ID'], int(r['Batch_number'])) for r in data]
expected = {(s, b) for s in 'ABCDEFG' for b in [1, 2, 3]}
require(len(data) == 21 and len(set(keys)) == 21 and (set(keys) == expected), 'Validation failed: len(data) == 21 and len(set(keys)) == 21 and (set(keys) == expected)')
require(len({r['Batch_ID'] for r in data}) == 21, "Validation failed: len({r['Batch_ID'] for r in data}) == 21")
form = {'A': ('Pangas', 50), 'B': ('Pangas', 40), 'C': ('Pangas', 30), 'D': ('Tilapia', 50), 'E': ('Tilapia', 40), 'F': ('Tilapia', 30), 'G': ('None', 0)}
for r in data:
    require((r['Fish_species'], r['Fish_inclusion_pct']) == form[r['Sample_ID']], "Validation failed: (r['Fish_species'], r['Fish_inclusion_pct']) == form[r['Sample_ID']]")
    require(r['Batch_ID'] == f"{r['Sample_ID']}_B{r['Batch_number']}", 'Validation failed: r[\'Batch_ID\'] == f"{r[\'Sample_ID\']}_B{r[\'Batch_number\']}"')
    require(r['Is_control'] == ('Yes' if r['Sample_ID'] == 'G' else 'No'), "Validation failed: r['Is_control'] == ('Yes' if r['Sample_ID'] == 'G' else 'No')")
a = np.array([[r[v] for v in VARIABLES] for r in data], float)
require(a.shape == (21, 10) and np.isfinite(a).all(), 'Validation failed: a.shape == (21, 10) and np.isfinite(a).all()')
require((np.std(a, axis=0, ddof=1) > 0).all(), 'Validation failed: (np.std(a, axis=0, ddof=1) > 0).all()')
checks = {}
for filename, cols in [('Fish_Crackers_Proximate_Preprocessed.xlsx', VARIABLES[:4]), ('Fish_Crackers_Physical_Color_Preprocessed.xlsx', VARIABLES[4:])]:
    w = openpyxl.load_workbook(DATA_FILE.parent / filename, data_only=True)
    source_rows = worksheet_records(w['Batch_data'], ['Sample_ID', 'Batch_number', *cols])
    lookup = keyed(source_rows, ['Sample_ID', 'Batch_number'])
    require(set(lookup) == set(keys), f'{filename}: batch keys differ from main workbook')
    n = 0
    for row, key in zip(data, keys):
        for col in cols:
            require(math.isclose(row[col], lookup[key][col], rel_tol=0, abs_tol=1e-12), (key, col, row[col], lookup[key][col]))
            n += 1
    checks[filename] = n
carb = np.array([r['Carbohydrate_pct'] for r in data])
closure = np.sum(a[:, :4], axis=1) + carb
require(np.allclose(closure, 100, atol=1e-08, rtol=0), 'Carbohydrate closure differs from 100%')
z_records = worksheet_records(wb['PCA_scaled'], ['Sample_ID', 'Batch_number', *[v + '_z' for v in VARIABLES]])
z_lookup = keyed(z_records, ['Sample_ID', 'Batch_number'])
require(set(z_lookup) == set(keys), 'PCA_scaled batch keys do not match')
z = np.array([[z_lookup[k][v + '_z'] for v in VARIABLES] for k in keys], float)
require(np.allclose(z, (a - a.mean(0)) / a.std(0, ddof=1), atol=1e-12, rtol=1e-10), 'PCA_scaled values mismatch')
means = np.array([a[[r['Sample_ID'] == sample for r in data]].mean(0) for sample in 'ABCDEFG'])
h_lookup = keyed(worksheet_records(wb['HCA_means'], ['Sample_ID', *VARIABLES]), ['Sample_ID'])
require(set(h_lookup) == {(sample,) for sample in 'ABCDEFG'}, 'HCA_means labels mismatch')
h_means = np.array([[h_lookup[sample,][v] for v in VARIABLES] for sample in 'ABCDEFG'], float)
require(np.allclose(means, h_means, atol=1e-12, rtol=1e-10), 'HCA_means values mismatch')

def correlations(matrix, method):
    return calculate_correlations(matrix, method, VARIABLES)
corr, p = correlations(a, 'Pearson')
scorr, s = correlations(a, 'Spearman')
keep = np.array([r['Sample_ID'] != 'G' for r in data])
ncorr, ng = correlations(a[keep], 'Pearson')
residual = a.copy()
for sample in 'ABCDEFG':
    ix = np.array([r['Sample_ID'] == sample for r in data])
    residual[ix] -= a[ix].mean(0)
within = np.corrcoef(residual, rowvar=False)
for row, i, j in zip(p, *zip(*list(itertools.combinations(range(10), 2)))):
    row['within_formulation_descriptive_r'] = float(within[i, j])
report = {'audit': {'rows': len(data), 'variables': VARIABLES, 'unique_batches': len(set(keys)), 'missing': int(np.isnan(a).sum()), 'source_reconciled': checks, 'closure_max_error': float(abs(closure - 100).max()), 'duplicate_numeric_rows': 21 - len(np.unique(a, axis=0)), 'PCA_scaled_verified': True, 'HCA_means_verified': True, 'pairs_each': 45}, 'pearson': p, 'spearman': s, 'pearson_without_G': ng, 'raw': data, 'matrices': {'pearson': corr.tolist(), 'spearman': scorr.tolist(), 'without_G': ncorr.tolist()}, 'probability_notes': {'pearson': 'Two-sided beta null; exploratory pooled fixed-formulation associations', 'spearman': 't approximation with average ranks; not an exact permutation test', 'BH': '45 unique pairs separately per analysis'}, 'significant_counts': {'pearson': sum((r['q'] < 0.05 for r in p)), 'spearman_approx': sum((r['q'] < 0.05 for r in s)), 'without_G': sum((r['q'] < 0.05 for r in ng))}}
import io, base64
from reportlab.pdfgen import canvas
from reportlab.lib.colors import Color, HexColor
import pypdfium2 as pdfium
from PIL import Image
short = ['Protein', 'Ash', 'Fat', 'Moisture', 'Expansion', 'Oil index', 'Specific vol.', 'L*', 'a*', 'b*']
units = ['%', '%', '%', '%', '%', '%', 'cm3/g', '', '', '']

def diverge(v):
    end = np.array([0.12, 0.34, 0.6]) if v < 0 else np.array([0.73, 0.27, 0.05])
    return Color(*(1 - abs(v)) * np.array([0.98, 0.98, 0.97]) + abs(v) * end)

def as_image(pdfbytes, scale=1.5, index=0):
    doc = pdfium.PdfDocument(pdfbytes)
    page = doc[index]
    pil = page.render(scale=scale).to_pil()
    buf = io.BytesIO()
    pil.save(buf, format='PNG')
    page.close()
    doc.close()
    return base64.b64encode(buf.getvalue()).decode()
FIGURES = OUTPUT_DIR / 'figures'
FIGURES.mkdir(parents=True, exist_ok=True)

def save_pdf_pages(pdfbytes, combined_name, page_names):
    """Export a vector multipage PDF and individual vector PDF/350-dpi PNG pages."""
    (FIGURES / (combined_name + '.pdf')).write_bytes(pdfbytes)
    document = pdfium.PdfDocument(pdfbytes)
    require(len(document) == len(page_names), 'Validation failed: len(document) == len(page_names)')
    for index, name in enumerate(page_names):
        single = pdfium.PdfDocument.new()
        single.import_pages(document, [index])
        single.save(str(FIGURES / (name + '.pdf')))
        single.close()
        page = document[index]
        bitmap = page.render(scale=FIGURE_DPI / 72)
        pil = bitmap.to_pil()
        pil.save(FIGURES / (name + '.png'), dpi=(FIGURE_DPI, FIGURE_DPI))
        bitmap.close()
        page.close()
    document.close()
heatpdf = io.BytesIO()
can = canvas.Canvas(heatpdf, pagesize=(690, 690))
for title, mat, n in [('Pearson correlations: all formulations', corr, 21), ('Spearman correlations: all formulations', scorr, 21), ('Pearson correlations: excluding G', ncorr, 18)]:
    can.setFillColor(HexColor('#202A35'))
    can.setFont('Helvetica-Bold', 16)
    can.drawString(35, 650, title)
    can.setFont('Helvetica', 10)
    can.drawString(35, 628, f'{n} independent batches; carbohydrate excluded; coefficients shown without significance stars')
    x0, y0, size = (128, 115, 45)
    for i in range(10):
        can.setFillColor(HexColor('#202A35'))
        can.setFont('Helvetica', 10)
        can.drawRightString(x0 - 8, y0 + (9 - i) * size + 19, short[i])
        can.saveState()
        can.translate(x0 + i * size + 20, y0 - 9)
        can.rotate(55)
        can.drawRightString(0, 0, short[i])
        can.restoreState()
        for j in range(10):
            v = float(mat[i, j])
            can.setFillColor(diverge(v))
            can.rect(x0 + j * size, y0 + (9 - i) * size, size, size, fill=1, stroke=0)
            can.setFillColor(Color(1, 1, 1) if abs(v) > 0.65 else HexColor('#202A35'))
            can.setFont('Helvetica', 9)
            can.drawCentredString(x0 + j * size + size / 2, y0 + (9 - i) * size + 19, f'{(0.0 if abs(v) < 0.005 else v):.2f}')
    for k in range(200):
        can.setFillColor(diverge(-1 + 2 * k / 199))
        can.rect(602, 115 + k * 2.25, 15, 2.5, fill=1, stroke=0)
    can.setFillColor(HexColor('#202A35'))
    can.setFont('Helvetica', 10)
    for val in [-1, -0.5, 0, 0.5, 1]:
        can.drawString(623, 115 + (val + 1) * 225 - 3, str(val))
    can.setFont('Helvetica', 9)
    can.drawString(35, 24, 'Pooled associations can reflect formulation differences; oil index is based on frying weight change.')
    can.showPage()
can.save()
save_pdf_pages(heatpdf.getvalue(), 'heatmaps_all', ['pearson_all_formulations', 'spearman_all_formulations', 'pearson_excluding_control_G'])
palette = ['#0072B2', '#E69F00', '#009E73', '#D55E00', '#CC79A7', '#56B4E9', '#222222']

def marker(can, x, y, k):
    can.setStrokeColor(HexColor(palette[k]))
    can.setFillColor(HexColor(palette[k]))
    r = 3
    if k == 0:
        can.circle(x, y, r, fill=1, stroke=0)
    elif k == 1:
        can.rect(x - r, y - r, 2 * r, 2 * r, fill=1, stroke=0)
    elif k in [2, 3, 4]:
        count = 3 if k in [2, 3] else 4
        offset = math.pi / 2 if k != 3 else -math.pi / 2
        path = can.beginPath()
        for j in range(count):
            xx = x + r * 1.4 * math.cos(offset + j * 2 * math.pi / count)
            yy = y + r * 1.4 * math.sin(offset + j * 2 * math.pi / count)
            if j == 0:
                path.moveTo(xx, yy)
            else:
                path.lineTo(xx, yy)
        path.close()
        can.drawPath(path, fill=1, stroke=0)
    elif k == 5:
        can.line(x - r, y - r, x + r, y + r)
        can.line(x - r, y + r, x + r, y - r)
    else:
        can.circle(x, y, r, fill=0, stroke=1)
pairs = list(itertools.combinations(range(10), 2))
scatter = io.BytesIO()
can = canvas.Canvas(scatter, pagesize=(840, 870))
for page in range(5):
    can.setFillColor(HexColor('#202A35'))
    can.setFont('Helvetica-Bold', 16)
    can.drawString(28, 836, f'Batch-level scatterplots ({page * 9 + 1}-{page * 9 + 9} of 45)')
    can.setFont('Helvetica', 10)
    for k, sample in enumerate('ABCDEFG'):
        marker(can, 40 + k * 110, 814, k)
        can.setFillColor(HexColor('#202A35'))
        can.drawString(49 + k * 110, 811, f'{sample} (n=3)')
    for slot, (i, j) in enumerate(pairs[page * 9:page * 9 + 9]):
        col = slot % 3
        row = slot // 3
        x0 = 52 + col * 273
        y0 = 605 - row * 258
        ww = 205
        hh = 160
        x = a[:, i]
        y = a[:, j]
        xmin = x.min() - 0.07 * np.ptp(x)
        xmax = x.max() + 0.07 * np.ptp(x)
        ymin = y.min() - 0.07 * np.ptp(y)
        ymax = y.max() + 0.07 * np.ptp(y)
        can.setFillColor(HexColor('#202A35'))
        can.setFont('Helvetica-Bold', 9)
        can.drawString(x0, y0 + hh + 16, f'{short[i]} vs {short[j]}  (r={corr[i, j]:.2f})')
        can.setStrokeColor(HexColor('#808080'))
        can.setLineWidth(0.5)
        can.line(x0, y0, x0 + ww, y0)
        can.line(x0, y0, x0, y0 + hh)
        can.setFont('Helvetica', 8)
        for f in [0, 0.5, 1]:
            can.line(x0 + f * ww, y0, x0 + f * ww, y0 - 3)
            can.drawCentredString(x0 + f * ww, y0 - 14, f'{xmin + f * (xmax - xmin):.2g}')
            can.line(x0 - 3, y0 + f * hh, x0, y0 + f * hh)
            can.drawRightString(x0 - 6, y0 + f * hh - 3, f'{ymin + f * (ymax - ymin):.2g}')
        can.setFont('Helvetica', 9)
        can.drawCentredString(x0 + ww / 2, y0 - 31, short[i] + (' (' + units[i] + ')' if units[i] else ''))
        can.saveState()
        can.translate(x0 - 35, y0 + hh / 2)
        can.rotate(90)
        can.drawCentredString(0, 0, short[j] + (' (' + units[j] + ')' if units[j] else ''))
        can.restoreState()
        for k, (xx, yy) in enumerate(zip(x, y)):
            marker(can, x0 + (xx - xmin) / (xmax - xmin) * ww, y0 + (yy - ymin) / (ymax - ymin) * hh, 'ABCDEFG'.index(data[k]['Sample_ID']))
    can.setFillColor(HexColor('#202A35'))
    can.setFont('Helvetica', 9)
    can.drawString(28, 24, 'One point per independent batch. Colours and shapes identify formulations; G is the fish-free control.')
    can.showPage()
can.save()
save_pdf_pages(scatter.getvalue(), 'scatterplots_all_45_pairs', [f'scatterplots_page_{i + 1:02d}' for i in range(5)])
import csv, importlib.metadata
for name, table in [('pearson_all', p), ('spearman_all', s), ('pearson_without_G', ng)]:
    with (OUTPUT_DIR / f'{name}.csv').open('w', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(table[0]))
        writer.writeheader()
        writer.writerows(table)
(OUTPUT_DIR / 'analysis_results.json').write_text(json.dumps(report, indent=2))
versions = {name: importlib.metadata.version(name) for name in ['numpy', 'scipy', 'openpyxl', 'reportlab', 'pypdfium2', 'Pillow']}
(OUTPUT_DIR / 'package_versions.json').write_text(json.dumps(versions, indent=2))
for path in FIGURES.glob('*.png'):
    with Image.open(path) as im:
        require(im.info['dpi'][0] >= 349, "Validation failed: im.info['dpi'][0] >= 349")
        im.verify()
for name, count in [('heatmaps_all.pdf', 3), ('scatterplots_all_45_pairs.pdf', 5)]:
    doc = pdfium.PdfDocument(str(FIGURES / name))
    require(len(doc) == count, 'Validation failed: len(doc) == count')
    doc.close()
require(len(list(FIGURES.glob('*.png'))) == 8, "Validation failed: len(list(FIGURES.glob('*.png'))) == 8")
print(json.dumps({'figures': str(FIGURES), 'png_files': 8, 'pdf_files': 10, 'dpi': FIGURE_DPI, 'pairs_per_analysis': 45}))

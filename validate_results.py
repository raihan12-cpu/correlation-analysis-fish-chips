# ============================================================
# USER SETTINGS / EDIT HERE
# ============================================================
from pathlib import Path
from correlation_core import require
ROOT = Path(__file__).resolve().parent
TOLERANCE = 1e-11
import csv, json, math, statistics, hashlib
import numpy as np
import openpyxl
from scipy.special import fdtrc
from PIL import Image
import pypdfium2 as pdfium

def average_ranks(values):
    """Assign average ranks to tied observations."""
    values = np.asarray(values)
    return np.array([1 + np.sum(values < x) + (np.sum(values == x) - 1) / 2 for x in values], float)

def main():
    result = json.loads((ROOT / 'results/analysis_results.json').read_text())
    wb = openpyxl.load_workbook(ROOT / 'source/Fish_Crackers_Correlation_PCA_HCA_Preprocessed.xlsx', data_only=True)
    rows = list(wb['Batch_data'].values)
    raw = [dict(zip(rows[0][:17], r[:17])) for r in rows[1:] if r[0]]
    variables = result['audit']['variables']
    checked = 0
    for key in ['pearson', 'spearman', 'pearson_without_G']:
        observations = [r for r in raw if key != 'pearson_without_G' or r['Sample_ID'] != 'G']
        pairs = result[key]
        require(len(pairs) == 45, 'Validation failed: len(pairs) == 45')
        sorted_p = sorted((r['p'] for r in pairs))
        for row in pairs:
            x = [r[row['x']] for r in observations]
            y = [r[row['y']] for r in observations]
            if key == 'spearman':
                x = average_ranks(x)
                y = average_ranks(y)
            corr = statistics.correlation(x, y)
            require(math.isclose(corr, row['r'], abs_tol=TOLERANCE), "Validation failed: math.isclose(corr, row['r'], abs_tol=TOLERANCE)")
            require(row['n'] == len(observations), "Validation failed: row['n'] == len(observations)")
            probability = 0.0 if abs(corr) >= 1 else float(fdtrc(1, len(x) - 2, corr * corr * (len(x) - 2) / (1 - corr * corr)))
            require(math.isclose(probability, row['p'], rel_tol=1e-09, abs_tol=1e-14), "Validation failed: math.isclose(probability, row['p'], rel_tol=1e-09, abs_tol=1e-14)")
            rank = sorted_p.index(row['p']) + 1
            adjusted = min(1, min((45 * sorted_p[j - 1] / j for j in range(rank, 46))))
            require(math.isclose(adjusted, row['q'], abs_tol=TOLERANCE), "Validation failed: math.isclose(adjusted, row['q'], abs_tol=TOLERANCE)")
            checked += 1
    excel_cells = 0
    if (ROOT / 'results/correlation_results.xlsx').exists():
        out = openpyxl.load_workbook(ROOT / 'results/correlation_results.xlsx', data_only=True)
        for sheet, key in [('Pearson A-G', 'pearson'), ('Spearman A-G', 'spearman'), ('Pearson A-F', 'pearson_without_G')]:
            for i, row in enumerate(result[key], 5):
                for j, k in [(3, 'n'), (4, 'r'), (5, 'p'), (6, 'q')]:
                    require(math.isclose(out[sheet].cell(i, j).value, row[k], rel_tol=1e-12, abs_tol=1e-14), 'Validation failed: math.isclose(out[sheet].cell(i, j).value, row[k], rel_tol=1e-12, abs_tol=1e-14)')
        excel_cells = 540
    figure_dir = ROOT / 'results/figures'
    pngs = list(figure_dir.glob('*.png'))
    require(len(pngs) == 8, 'Validation failed: len(pngs) == 8')
    for path in pngs:
        with Image.open(path) as im:
            require(min(im.size) > 3000 and im.info['dpi'][0] >= 349, "Validation failed: min(im.size) > 3000 and im.info['dpi'][0] >= 349")
            im.verify()
    for filename, count in [('heatmaps_all.pdf', 3), ('scatterplots_all_45_pairs.pdf', 5)]:
        document = pdfium.PdfDocument(str(figure_dir / filename))
        require(len(document) == count, 'Validation failed: len(document) == count')
        document.close()
    record = {'coefficient_p_and_BH_checks': checked, 'excel_numeric_cells_checked': excel_cells, 'png_files_checked': 8, 'heatmap_pages': 3, 'scatterplot_pages': 5, 'source_sha256': hashlib.sha256((ROOT / 'source/Fish_Crackers_Correlation_PCA_HCA_Preprocessed.xlsx').read_bytes()).hexdigest()}
    (ROOT / 'results/validation.json').write_text(json.dumps(record, indent=2))
    print(json.dumps(record))
if __name__ == '__main__':
    main()
